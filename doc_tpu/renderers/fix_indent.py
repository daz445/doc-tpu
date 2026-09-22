"""Заменяет firstLine indent на tab- indent в абзацах тела отчёта.

1. Убирает w:firstLine из XML
2. Ставит tab stop (w:tabs) на 1.2 cm
3. Вставляет \t в начало первого run
"""

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm
from pathlib import Path
import sys

INDENT_CM = 1.2
# 1 cm = 360000 EMU, 1 twip = 635 EMU → 1 cm ≈ 567 twips
INDENT_TWIPS = int(INDENT_CM * 567)


def add_tab_to_paragraphs(doc_path: str):
    doc = Document(doc_path)
    after_page_break = False

    for para in doc.paragraphs:
        # Пропускаем титульную страницу
        if not after_page_break:
            for run in para.runs:
                if run._r.findall(qn("w:br")):
                    after_page_break = True
                    break
            pPr = para._element.find(qn("w:pPr"))
            if pPr is not None:
                for child in pPr:
                    if child.tag == qn("w:pageBreakBefore"):
                        after_page_break = True
            if not after_page_break:
                continue

        # Пропускаем пустые абзацы
        if not para.text.strip():
            continue

        pPr = para._element.get_or_add_pPr()

        # 1. Убираем firstLine indent
        ind = pPr.find(qn("w:ind"))
        if ind is not None:
            if ind.get(qn("w:firstLine")) is not None:
                del ind.attrib[qn("w:firstLine")]
            if len(ind.attrib) == 0:
                pPr.remove(ind)

        # 2. Ставим tab stop на 1.2 cm (left align)
        tabs_elem = pPr.find(qn("w:tabs"))
        if tabs_elem is None:
            tabs_elem = OxmlElement("w:tabs")
            pPr.append(tabs_elem)

        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "left")
        tab.set(qn("w:pos"), str(INDENT_TWIPS))
        tabs_elem.append(tab)

        # 3. Вставляем \t в начало первого run
        if para.runs:
            first_run = para.runs[0]
            first_run.text = "\t" + first_run.text
        else:
            run = para.add_run("\t")

    doc.save(doc_path)
    print(f"Исправлено: {doc_path}")


if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else str(
        Path(__file__).parent / "report.docx"
    )
    add_tab_to_paragraphs(path)
